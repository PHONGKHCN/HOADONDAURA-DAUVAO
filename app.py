import os
import json
import re
import sqlite3
from datetime import datetime
from io import BytesIO

from flask import Flask, render_template, request, redirect, url_for, flash, send_file, g, session
from flask_login import (
    LoginManager, UserMixin, login_user, logout_user,
    login_required, current_user
)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import A4, A5
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from vn_number_to_words import so_thanh_chu

from docx import Document as DocxDocument
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DB_PATH = os.environ.get("DB_PATH") or os.path.join(os.path.dirname(__file__), "hoadon.db")

FONT_DIR = os.path.join(os.path.dirname(__file__), "static", "fonts")
pdfmetrics.registerFont(TTFont("VNSans", os.path.join(FONT_DIR, "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("VNSans-Bold", os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf")))

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "doi-secret-key-nay-khi-deploy")

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"
login_manager.login_message = "Vui lòng đăng nhập để tiếp tục."


# ---------- DB helpers ----------

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
        g.db.execute("PRAGMA journal_mode = WAL")
        g.db.execute("PRAGMA busy_timeout = 5000")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            full_name TEXT,
            role TEXT NOT NULL DEFAULT 'staff',  -- 'admin' | 'staff'
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            loai TEXT NOT NULL,              -- 'dau_ra' | 'dau_vao'
            nhom TEXT,                       -- nhóm/khách hàng (VLXD, tạp hóa, v.v.)
            ky_hieu TEXT,
            so_hd TEXT NOT NULL,
            ngay_lap TEXT NOT NULL,          -- YYYY-MM-DD
            doi_tac TEXT NOT NULL,
            mst TEXT,
            mat_hang TEXT,
            doanh_so REAL NOT NULL DEFAULT 0,
            thue_suat REAL NOT NULL DEFAULT 10,
            tien_thue REAL NOT NULL DEFAULT 0,
            tong_cong REAL NOT NULL DEFAULT 0,
            ghi_chu TEXT,
            nguoi_tao_id INTEGER,
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (nguoi_tao_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS sales_invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nhom TEXT,                    -- nhóm/khách hàng (VLXD, tạp hóa, v.v.)
            so_hd TEXT,
            ngay_lap TEXT NOT NULL,
            seller_name TEXT NOT NULL,
            seller_address TEXT,
            seller_phone TEXT,
            buyer_name TEXT NOT NULL,
            buyer_address TEXT,
            tong_cong REAL NOT NULL DEFAULT 0,
            nguoi_tao_id INTEGER,
            created_at TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (nguoi_tao_id) REFERENCES users(id)
        );

        CREATE TABLE IF NOT EXISTS sales_invoice_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            invoice_id INTEGER NOT NULL,
            stt INTEGER NOT NULL,
            ten_hang TEXT NOT NULL,
            quy_cach TEXT,
            dvt TEXT,
            so_luong REAL NOT NULL DEFAULT 0,
            don_gia REAL NOT NULL DEFAULT 0,
            thanh_tien REAL NOT NULL DEFAULT 0,
            FOREIGN KEY (invoice_id) REFERENCES sales_invoices(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS partners (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            loai TEXT NOT NULL,           -- 'seller' | 'buyer'
            nhom TEXT,                    -- nhóm/khách hàng (vd: "Nguyen Van Do - VLXD")
            ten TEXT NOT NULL,
            dia_chi TEXT,
            sdt TEXT,
            mst TEXT,
            ghi_chu TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nhom TEXT,                    -- nhóm/khách hàng (để tách danh mục riêng cho từng khách)
            ten TEXT NOT NULL,
            quy_cach TEXT,
            dvt TEXT,
            gia_ban REAL DEFAULT 0,
            gia_mua REAL DEFAULT 0,
            ghi_chu TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );
        """
    )
    db.commit()

    # Migration an toàn: thêm cột 'nhom' nếu DB được tạo từ bản cũ chưa có cột này
    for table in ("invoices", "sales_invoices"):
        cols = [r["name"] for r in db.execute(f"PRAGMA table_info({table})").fetchall()]
        if "nhom" not in cols:
            db.execute(f"ALTER TABLE {table} ADD COLUMN nhom TEXT")
    db.commit()

    # Migration: thêm các cột mới cho mẫu "Hóa đơn bán lẻ" (số điện thoại người mua,
    # địa điểm lập hóa đơn, dòng mô tả nhỏ dưới tên bên bán)
    sales_cols = [r["name"] for r in db.execute("PRAGMA table_info(sales_invoices)").fetchall()]
    for col in ("buyer_phone", "dia_diem", "seller_slogan", "loai_chung_tu"):
        if col not in sales_cols:
            db.execute(f"ALTER TABLE sales_invoices ADD COLUMN {col} TEXT")
    if "thanh_toan" not in sales_cols:
        db.execute("ALTER TABLE sales_invoices ADD COLUMN thanh_toan REAL DEFAULT 0")
    db.commit()

    # Seed one admin account if no users exist yet
    cur = db.execute("SELECT COUNT(*) AS c FROM users")
    if cur.fetchone()["c"] == 0:
        db.execute(
            "INSERT INTO users (username, password_hash, full_name, role) VALUES (?, ?, ?, ?)",
            ("admin", generate_password_hash("admin123"), "Quản trị viên", "admin"),
        )
        db.commit()
        print("Đã tạo tài khoản mặc định: admin / admin123  (đổi mật khẩu ngay sau khi đăng nhập lần đầu)")
    db.close()


# ---------- User model ----------

class User(UserMixin):
    def __init__(self, row):
        self.id = row["id"]
        self.username = row["username"]
        self.full_name = row["full_name"]
        self.role = row["role"]

    @property
    def is_admin(self):
        return self.role == "admin"


@login_manager.user_loader
def load_user(user_id):
    db = get_db()
    row = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    return User(row) if row else None


def admin_required(view):
    from functools import wraps

    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash("Chỉ quản trị viên mới có quyền truy cập chức năng này.", "danger")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)

    return wrapped


# ---------- Auth routes ----------

@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        db = get_db()
        row = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        if row and check_password_hash(row["password_hash"], password):
            login_user(User(row))
            return redirect(url_for("dashboard"))
        flash("Sai tên đăng nhập hoặc mật khẩu.", "danger")
    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


# ---------- Dashboard / reconciliation ----------

def month_bounds(period):
    """period format YYYY-MM -> (start_date, end_date) as strings YYYY-MM-DD"""
    year, month = map(int, period.split("-"))
    start = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end = f"{year+1:04d}-01-01"
    else:
        end = f"{year:04d}-{month+1:02d}-01"
    return start, end


@app.route("/")
@login_required
def dashboard():
    period = request.args.get("period") or datetime.now().strftime("%Y-%m")
    start, end = month_bounds(period)
    db = get_db()

    def totals(loai):
        row = db.execute(
            """SELECT COALESCE(SUM(doanh_so),0) AS doanh_so,
                      COALESCE(SUM(tien_thue),0) AS tien_thue,
                      COALESCE(SUM(tong_cong),0) AS tong_cong,
                      COUNT(*) AS so_luong
               FROM invoices
               WHERE loai = ? AND ngay_lap >= ? AND ngay_lap < ?""",
            (loai, start, end),
        ).fetchone()
        return row

    dau_ra = totals("dau_ra")
    dau_vao = totals("dau_vao")
    chenh_lech = dau_ra["tong_cong"] - dau_vao["tong_cong"]

    return render_template(
        "dashboard.html",
        period=period,
        dau_ra=dau_ra,
        dau_vao=dau_vao,
        chenh_lech=chenh_lech,
    )


# ---------- Invoice CRUD ----------

@app.route("/invoices/bulk", methods=["GET", "POST"])
@login_required
def invoice_bulk():
    if request.method == "POST":
        raw = request.form.get("bulk_text", "")
        default_loai = request.form.get("default_loai", "")
        default_nhom = request.form.get("default_nhom", "").strip()
        lines = [l for l in raw.replace("\r\n", "\n").split("\n") if l.strip()]

        db = get_db()
        created = 0
        errors = []
        for idx, line in enumerate(lines, start=1):
            cols = [c.strip() for c in line.split("\t")]
            if len(cols) < 4:
                errors.append(f"Dòng {idx}: thiếu cột (cần ít nhất 4 cột: Loại, Ngày, Đối tác, Số tiền)")
                continue
            try:
                loai_raw = cols[0].lower()
                if loai_raw in ("dau_vao", "đầu vào", "vao", "mua", "vào"):
                    loai = "dau_vao"
                elif loai_raw in ("dau_ra", "đầu ra", "ra", "ban", "bán"):
                    loai = "dau_ra"
                else:
                    loai = default_loai if default_loai in ("dau_vao", "dau_ra") else None
                if not loai:
                    errors.append(f"Dòng {idx}: không nhận diện được Loại ('{cols[0]}')")
                    continue

                ngay_raw = cols[1].strip()
                ngay_lap = None
                for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                    try:
                        ngay_lap = datetime.strptime(ngay_raw, fmt).strftime("%Y-%m-%d")
                        break
                    except ValueError:
                        continue
                if not ngay_lap:
                    errors.append(f"Dòng {idx}: ngày không hợp lệ ('{ngay_raw}')")
                    continue

                doi_tac = cols[2]
                mat_hang = cols[3] if len(cols) > 3 else ""
                so_tien_raw = cols[4] if len(cols) > 4 else "0"
                doanh_so = float(re.sub(r"[^\d.\-]", "", so_tien_raw.replace(",", "")) or 0)
                so_hd = cols[5] if len(cols) > 5 else ""
                ky_hieu = cols[6] if len(cols) > 6 else ""
                ghi_chu = cols[7] if len(cols) > 7 else ""

                if not doi_tac:
                    errors.append(f"Dòng {idx}: thiếu tên đối tác")
                    continue

                db.execute(
                    """INSERT INTO invoices
                       (loai, nhom, ky_hieu, so_hd, ngay_lap, doi_tac, mst, mat_hang,
                        doanh_so, thue_suat, tien_thue, tong_cong, ghi_chu, nguoi_tao_id)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (loai, default_nhom, ky_hieu, so_hd, ngay_lap, doi_tac, "", mat_hang,
                     doanh_so, 0, 0, doanh_so, ghi_chu, current_user.id),
                )
                created += 1
            except Exception as e:
                errors.append(f"Dòng {idx}: lỗi xử lý ({e})")

        db.commit()
        if created:
            flash(f"Đã tạo {created} hóa đơn.", "success")
        if errors:
            flash("Một số dòng bị bỏ qua: " + " | ".join(errors[:10]) +
                  (f" (và {len(errors)-10} lỗi khác)" if len(errors) > 10 else ""), "danger")
        return redirect(url_for("invoice_list"))

    return render_template("invoice_bulk.html")


@app.route("/invoices")
@login_required
def invoice_list():
    loai = request.args.get("loai", "")
    period = request.args.get("period", "")
    nhom = request.args.get("nhom", "")
    q = request.args.get("q", "").strip()

    sql = """SELECT inv.*, u.full_name AS nguoi_tao_ten, u.username AS nguoi_tao_username
             FROM invoices inv LEFT JOIN users u ON inv.nguoi_tao_id = u.id WHERE 1=1"""
    params = []
    if loai in ("dau_ra", "dau_vao"):
        sql += " AND inv.loai = ?"
        params.append(loai)
    if nhom:
        sql += " AND inv.nhom = ?"
        params.append(nhom)
    if period:
        start, end = month_bounds(period)
        sql += " AND inv.ngay_lap >= ? AND inv.ngay_lap < ?"
        params += [start, end]
    if q:
        sql += " AND (inv.doi_tac LIKE ? OR inv.so_hd LIKE ?)"
        like = f"%{q}%"
        params += [like, like]
    sql += " ORDER BY inv.ngay_lap DESC, inv.id DESC"

    db = get_db()
    invoices = db.execute(sql, params).fetchall()
    nhoms = db.execute(
        "SELECT DISTINCT nhom FROM invoices WHERE nhom IS NOT NULL AND nhom != '' ORDER BY nhom"
    ).fetchall()
    return render_template(
        "invoice_list.html", invoices=invoices, loai=loai, period=period, q=q,
        nhom=nhom, nhoms=nhoms,
    )


def parse_invoice_form(form):
    doanh_so = float(form.get("doanh_so") or 0)
    return {
        "loai": form.get("loai"),
        "nhom": form.get("nhom", "").strip(),
        "ky_hieu": form.get("ky_hieu", "").strip(),
        "so_hd": form.get("so_hd", "").strip(),
        "ngay_lap": form.get("ngay_lap"),
        "doi_tac": form.get("doi_tac", "").strip(),
        "mat_hang": form.get("mat_hang", "").strip(),
        "doanh_so": doanh_so,
        "ghi_chu": form.get("ghi_chu", "").strip(),
    }


@app.route("/invoices/new", methods=["GET", "POST"])
@login_required
def invoice_new():
    if request.method == "POST":
        data = parse_invoice_form(request.form)
        db = get_db()
        db.execute(
            """INSERT INTO invoices
               (loai, nhom, ky_hieu, so_hd, ngay_lap, doi_tac, mst, mat_hang,
                doanh_so, thue_suat, tien_thue, tong_cong, ghi_chu, nguoi_tao_id)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                data["loai"], data["nhom"], data["ky_hieu"], data["so_hd"], data["ngay_lap"],
                data["doi_tac"], "", data["mat_hang"], data["doanh_so"],
                0, 0, data["doanh_so"], data["ghi_chu"], current_user.id,
            ),
        )
        db.commit()
        flash("Đã lưu hóa đơn.", "success")
        return redirect(url_for("invoice_list", loai=data["loai"]))
    return render_template("invoice_form.html", invoice=None)


@app.route("/invoices/<int:invoice_id>/edit", methods=["GET", "POST"])
@login_required
def invoice_edit(invoice_id):
    db = get_db()
    invoice = db.execute("SELECT * FROM invoices WHERE id = ?", (invoice_id,)).fetchone()
    if not invoice:
        flash("Không tìm thấy hóa đơn.", "danger")
        return redirect(url_for("invoice_list"))

    if request.method == "POST":
        data = parse_invoice_form(request.form)
        db.execute(
            """UPDATE invoices SET loai=?, nhom=?, ky_hieu=?, so_hd=?, ngay_lap=?, doi_tac=?,
               mat_hang=?, doanh_so=?, tong_cong=?, ghi_chu=?
               WHERE id=?""",
            (
                data["loai"], data["nhom"], data["ky_hieu"], data["so_hd"], data["ngay_lap"],
                data["doi_tac"], data["mat_hang"], data["doanh_so"], data["doanh_so"],
                data["ghi_chu"], invoice_id,
            ),
        )
        db.commit()
        flash("Đã cập nhật hóa đơn.", "success")
        return redirect(url_for("invoice_list", loai=data["loai"]))

    return render_template("invoice_form.html", invoice=invoice)


@app.route("/invoices/<int:invoice_id>/delete", methods=["POST"])
@login_required
def invoice_delete(invoice_id):
    db = get_db()
    db.execute("DELETE FROM invoices WHERE id = ?", (invoice_id,))
    db.commit()
    flash("Đã xóa hóa đơn.", "success")
    return redirect(request.referrer or url_for("invoice_list"))


# ---------- Excel export (bảng kê đầu vào / đầu ra) ----------

FONT_NAME = "Times New Roman"


def style_ledger_sheet(ws, title, period, rows):
    header_fill = PatternFill(start_color="2E5266", end_color="2E5266", fill_type="solid")
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
    title_font = Font(name=FONT_NAME, bold=True, size=13)
    sub_font = Font(name=FONT_NAME, italic=True, size=10)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    headers = ["STT", "Ký hiệu HĐ", "Số HĐ", "Ngày lập", "Tên đối tác",
               "Mặt hàng/Dịch vụ", "Số tiền (đ)", "Ghi chú"]
    widths = [5, 10, 8, 11, 24, 26, 16, 20]

    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Kỳ: {period}"
    ws["A2"].font = sub_font
    ws["A2"].alignment = center

    hr = 4
    for i, h in enumerate(headers):
        c = ws.cell(row=hr, column=1 + i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = center

    r = hr + 1
    for idx, inv in enumerate(rows, start=1):
        values = [
            idx, inv["ky_hieu"], inv["so_hd"], inv["ngay_lap"], inv["doi_tac"],
            inv["mat_hang"], inv["tong_cong"], inv["ghi_chu"],
        ]
        for i, val in enumerate(values):
            c = ws.cell(row=r, column=1 + i, value=val)
            c.border = border
            c.font = Font(name=FONT_NAME, size=10)
            if i in (0, 1, 2, 3):
                c.alignment = center
            elif i == 6:
                c.alignment = right
                c.number_format = "#,##0"
            else:
                c.alignment = left
        r += 1

    if not rows:
        r += 1  # avoid SUM over header only

    total_row = r
    ws.cell(row=total_row, column=6, value="TỔNG CỘNG").font = Font(name=FONT_NAME, bold=True, size=10)
    ws.cell(row=total_row, column=6).alignment = Alignment(horizontal="right")
    cell = ws.cell(row=total_row, column=7, value=f"=SUM(G{hr+1}:G{total_row-1})")
    cell.font = Font(name=FONT_NAME, bold=True, size=10)
    cell.number_format = "#,##0"
    cell.alignment = right
    for col_idx in range(1, 9):
        ws.cell(row=total_row, column=col_idx).border = border
        if col_idx != 7:
            ws.cell(row=total_row, column=col_idx).fill = PatternFill(
                start_color="D6E4F0", end_color="D6E4F0", fill_type="solid"
            )

    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A5"
    return total_row


@app.route("/export")
@login_required
def export_excel():
    period = request.args.get("period") or datetime.now().strftime("%Y-%m")
    start, end = month_bounds(period)
    db = get_db()

    dau_ra = db.execute(
        "SELECT * FROM invoices WHERE loai='dau_ra' AND ngay_lap >= ? AND ngay_lap < ? ORDER BY ngay_lap",
        (start, end),
    ).fetchall()
    dau_vao = db.execute(
        "SELECT * FROM invoices WHERE loai='dau_vao' AND ngay_lap >= ? AND ngay_lap < ? ORDER BY ngay_lap",
        (start, end),
    ).fetchall()

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Đầu ra (Bán ra)"
    r1 = style_ledger_sheet(ws1, "BẢNG KÊ HÓA ĐƠN - ĐẦU RA (BÁN RA)", period, dau_ra)

    ws2 = wb.create_sheet("Đầu vào (Mua vào)")
    r2 = style_ledger_sheet(ws2, "BẢNG KÊ HÓA ĐƠN - ĐẦU VÀO (MUA VÀO)", period, dau_vao)

    ws3 = wb.create_sheet("Đối chiếu")
    title_font = Font(name=FONT_NAME, bold=True, size=13)
    header_fill = PatternFill(start_color="2E5266", end_color="2E5266", fill_type="solid")
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center")
    left = Alignment(horizontal="left")
    right = Alignment(horizontal="right")

    ws3.merge_cells("A1:C1")
    ws3["A1"] = f"ĐỐI CHIẾU ĐẦU VÀO - ĐẦU RA (Kỳ {period})"
    ws3["A1"].font = title_font
    ws3["A1"].alignment = center

    labels = [
        ("Chỉ tiêu", "Đầu ra", "Đầu vào"),
        ("Tổng số tiền", f"='Đầu ra (Bán ra)'!G{r1}", f"='Đầu vào (Mua vào)'!G{r2}"),
    ]
    for i, row in enumerate(labels):
        for j, val in enumerate(row):
            c = ws3.cell(row=5 + i, column=1 + j, value=val)
            c.border = border
            if i == 0:
                c.font = header_font
                c.fill = header_fill
                c.alignment = center
            else:
                c.font = Font(name=FONT_NAME, size=10, bold=(j == 0))
                c.alignment = left if j == 0 else right
                if j > 0:
                    c.number_format = "#,##0"
    ws3.cell(row=7, column=1, value="Chênh lệch (Đầu ra - Đầu vào)").font = Font(name=FONT_NAME, size=10, bold=True)
    ws3.cell(row=7, column=1).border = border
    ws3.cell(row=7, column=1).alignment = left
    diff_cell = ws3.cell(row=7, column=2, value="=B6-C6")
    diff_cell.font = Font(name=FONT_NAME, size=10, bold=True)
    diff_cell.number_format = "#,##0"
    diff_cell.alignment = right
    diff_cell.border = border
    ws3.cell(row=7, column=3).border = border
    for col, w in zip("ABC", [28, 18, 18]):
        ws3.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Bang_ke_hoa_don_{period}.xlsx"
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------- User management (admin only) ----------

@app.route("/backup")
@login_required
@admin_required
def backup_download():
    """Tải file database hiện tại về máy để sao lưu thủ công."""
    if not os.path.exists(DB_PATH):
        flash("Chưa có dữ liệu để sao lưu.", "danger")
        return redirect(url_for("dashboard"))
    # Chống ghi đè khi đang có transaction dở dang: chép ra buffer trước khi gửi
    with sqlite3.connect(DB_PATH) as src:
        buf = BytesIO()
        for line in src.iterdump():
            buf.write((line + "\n").encode("utf-8"))
    buf.seek(0)
    filename = f"backup_hoadon_{datetime.now().strftime('%Y%m%d_%H%M%S')}.sql"
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/sql")


@app.route("/restore", methods=["GET", "POST"])
@login_required
@admin_required
def restore_upload():
    if request.method == "POST":
        file = request.files.get("backup_file")
        if not file or not file.filename:
            flash("Vui lòng chọn file sao lưu (.sql).", "danger")
            return redirect(url_for("restore_upload"))

        # Chỉ khôi phục dữ liệu nghiệp vụ theo mặc định — KHÔNG đụng vào bảng 'users'
        # để tránh tự khóa tài khoản admin đang đăng nhập. Chỉ khôi phục cả 'users'
        # khi người dùng chủ động tick chọn.
        allowed_tables = ["invoices", "sales_invoices", "sales_invoice_items", "partners", "products"]
        if request.form.get("restore_users"):
            allowed_tables.append("users")
        try:
            sql_text = file.stream.read().decode("utf-8")
            insert_statements = []
            tables_seen = set()
            for line in sql_text.splitlines():
                line = line.strip()
                m = re.match(r'INSERT INTO ["\']?(\w+)["\']?', line, re.IGNORECASE)
                if m and m.group(1) in allowed_tables:
                    insert_statements.append(line)
                    tables_seen.add(m.group(1))

            if not insert_statements:
                flash("File không chứa dữ liệu hợp lệ để khôi phục.", "danger")
                return redirect(url_for("restore_upload"))

            db = get_db()
            for table in tables_seen:
                db.execute(f"DELETE FROM {table}")
            restored = 0
            skipped = 0
            for stmt in insert_statements:
                try:
                    db.execute(stmt)
                    restored += 1
                except Exception:
                    skipped += 1
            db.commit()
            msg = f"Đã khôi phục {restored} dòng dữ liệu từ file sao lưu."
            if skipped:
                msg += f" ({skipped} dòng bị bỏ qua do lỗi định dạng.)"
            flash(msg, "success")

            if "users" in tables_seen:
                # Tài khoản vừa bị thay thế toàn bộ — đăng xuất để tránh phiên đăng nhập
                # hiện tại trỏ vào tài khoản không còn tồn tại/đổi mật khẩu.
                logout_user()
                flash("Đã khôi phục tài khoản đăng nhập — vui lòng đăng nhập lại.", "success")
                return redirect(url_for("login"))
        except Exception as e:
            flash(f"Không khôi phục được: {e}", "danger")
        return redirect(url_for("dashboard"))
    return render_template("restore.html")


@app.route("/users")
@login_required
@admin_required
def user_list():
    db = get_db()
    users = db.execute("SELECT * FROM users ORDER BY id").fetchall()
    return render_template("user_list.html", users=users)


@app.route("/users/new", methods=["GET", "POST"])
@login_required
@admin_required
def user_new():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        full_name = request.form.get("full_name", "").strip()
        role = request.form.get("role", "staff")
        db = get_db()
        exists = db.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()
        if exists:
            flash("Tên đăng nhập đã tồn tại.", "danger")
        else:
            db.execute(
                "INSERT INTO users (username, password_hash, full_name, role) VALUES (?,?,?,?)",
                (username, generate_password_hash(password), full_name, role),
            )
            db.commit()
            flash("Đã tạo tài khoản.", "success")
            return redirect(url_for("user_list"))
    return render_template("user_form.html")


@app.route("/users/<int:user_id>/delete", methods=["POST"])
@login_required
@admin_required
def user_delete(user_id):
    if user_id == current_user.id:
        flash("Không thể tự xóa tài khoản đang đăng nhập.", "danger")
        return redirect(url_for("user_list"))
    db = get_db()
    db.execute("DELETE FROM users WHERE id = ?", (user_id,))
    db.commit()
    flash("Đã xóa tài khoản.", "success")
    return redirect(url_for("user_list"))


@app.route("/account/change_password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        current_pw = request.form.get("current_password", "")
        new_pw = request.form.get("new_password", "")
        confirm_pw = request.form.get("confirm_password", "")

        db = get_db()
        row = db.execute("SELECT * FROM users WHERE id = ?", (current_user.id,)).fetchone()

        if not check_password_hash(row["password_hash"], current_pw):
            flash("Mật khẩu hiện tại không đúng.", "danger")
        elif len(new_pw) < 4:
            flash("Mật khẩu mới phải có ít nhất 4 ký tự.", "danger")
        elif new_pw != confirm_pw:
            flash("Mật khẩu mới nhập lại không khớp.", "danger")
        else:
            db.execute(
                "UPDATE users SET password_hash = ? WHERE id = ?",
                (generate_password_hash(new_pw), current_user.id),
            )
            db.commit()
            flash("Đã đổi mật khẩu thành công.", "success")
            return redirect(url_for("dashboard"))

    return render_template("change_password.html")


@app.route("/users/<int:user_id>/reset_password", methods=["GET", "POST"])
@login_required
@admin_required
def user_reset_password(user_id):
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if not user:
        flash("Không tìm thấy tài khoản.", "danger")
        return redirect(url_for("user_list"))

    if request.method == "POST":
        new_pw = request.form.get("new_password", "")
        confirm_pw = request.form.get("confirm_password", "")
        if len(new_pw) < 4:
            flash("Mật khẩu mới phải có ít nhất 4 ký tự.", "danger")
        elif new_pw != confirm_pw:
            flash("Mật khẩu nhập lại không khớp.", "danger")
        else:
            db.execute(
                "UPDATE users SET password_hash = ? WHERE id = ?",
                (generate_password_hash(new_pw), user_id),
            )
            db.commit()
            flash(f"Đã đặt lại mật khẩu cho tài khoản \"{user['username']}\".", "success")
            return redirect(url_for("user_list"))

    return render_template("reset_password.html", user=user)


# ---------- Danh mục đối tác (bên bán / bên mua đã lưu) ----------

@app.route("/partners")
@login_required
def partner_list():
    nhom = request.args.get("nhom", "")
    db = get_db()
    sql = "SELECT * FROM partners WHERE 1=1"
    params = []
    if nhom:
        sql += " AND nhom = ?"
        params.append(nhom)
    sql += " ORDER BY nhom, loai, ten"
    partners = db.execute(sql, params).fetchall()
    nhoms = db.execute(
        "SELECT DISTINCT nhom FROM partners WHERE nhom IS NOT NULL AND nhom != '' ORDER BY nhom"
    ).fetchall()
    return render_template("partner_list.html", partners=partners, nhoms=nhoms, nhom=nhom)


@app.route("/partners/new", methods=["GET", "POST"])
@login_required
def partner_new():
    if request.method == "POST":
        f = request.form
        db = get_db()
        db.execute(
            """INSERT INTO partners (loai, nhom, ten, dia_chi, sdt, mst, ghi_chu)
               VALUES (?,?,?,?,?,?,?)""",
            (f.get("loai"), f.get("nhom", "").strip(), f.get("ten", "").strip(),
             f.get("dia_chi", "").strip(), f.get("sdt", "").strip(),
             f.get("mst", "").strip(), f.get("ghi_chu", "").strip()),
        )
        db.commit()
        flash("Đã lưu đối tác.", "success")
        return redirect(url_for("partner_list"))
    return render_template("partner_form.html", partner=None)


@app.route("/partners/<int:partner_id>/edit", methods=["GET", "POST"])
@login_required
def partner_edit(partner_id):
    db = get_db()
    partner = db.execute("SELECT * FROM partners WHERE id=?", (partner_id,)).fetchone()
    if not partner:
        flash("Không tìm thấy đối tác.", "danger")
        return redirect(url_for("partner_list"))
    if request.method == "POST":
        f = request.form
        db.execute(
            """UPDATE partners SET loai=?, nhom=?, ten=?, dia_chi=?, sdt=?, mst=?, ghi_chu=?
               WHERE id=?""",
            (f.get("loai"), f.get("nhom", "").strip(), f.get("ten", "").strip(),
             f.get("dia_chi", "").strip(), f.get("sdt", "").strip(),
             f.get("mst", "").strip(), f.get("ghi_chu", "").strip(), partner_id),
        )
        db.commit()
        flash("Đã cập nhật đối tác.", "success")
        return redirect(url_for("partner_list"))
    return render_template("partner_form.html", partner=partner)


@app.route("/partners/<int:partner_id>/delete", methods=["POST"])
@login_required
def partner_delete(partner_id):
    db = get_db()
    db.execute("DELETE FROM partners WHERE id=?", (partner_id,))
    db.commit()
    flash("Đã xóa đối tác.", "success")
    return redirect(url_for("partner_list"))


# ---------- Danh mục sản phẩm (dùng để chọn nhanh khi tạo hóa đơn) ----------

@app.route("/products")
@login_required
def product_list():
    nhom = request.args.get("nhom", "")
    db = get_db()
    sql = "SELECT * FROM products WHERE 1=1"
    params = []
    if nhom:
        sql += " AND nhom = ?"
        params.append(nhom)
    sql += " ORDER BY nhom, ten"
    products = db.execute(sql, params).fetchall()
    nhoms = db.execute(
        "SELECT DISTINCT nhom FROM products WHERE nhom IS NOT NULL AND nhom != '' ORDER BY nhom"
    ).fetchall()
    return render_template("product_list.html", products=products, nhoms=nhoms, nhom=nhom)


@app.route("/products/new", methods=["GET", "POST"])
@login_required
def product_new():
    if request.method == "POST":
        f = request.form
        db = get_db()
        db.execute(
            """INSERT INTO products (nhom, ten, quy_cach, dvt, gia_ban, gia_mua, ghi_chu)
               VALUES (?,?,?,?,?,?,?)""",
            (f.get("nhom", "").strip(), f.get("ten", "").strip(), f.get("quy_cach", "").strip(),
             f.get("dvt", "").strip(), float(f.get("gia_ban") or 0), float(f.get("gia_mua") or 0),
             f.get("ghi_chu", "").strip()),
        )
        db.commit()
        flash("Đã lưu sản phẩm.", "success")
        return redirect(url_for("product_list"))
    return render_template("product_form.html", product=None)


@app.route("/products/<int:product_id>/edit", methods=["GET", "POST"])
@login_required
def product_edit(product_id):
    db = get_db()
    product = db.execute("SELECT * FROM products WHERE id=?", (product_id,)).fetchone()
    if not product:
        flash("Không tìm thấy sản phẩm.", "danger")
        return redirect(url_for("product_list"))
    if request.method == "POST":
        f = request.form
        db.execute(
            """UPDATE products SET nhom=?, ten=?, quy_cach=?, dvt=?, gia_ban=?, gia_mua=?, ghi_chu=?
               WHERE id=?""",
            (f.get("nhom", "").strip(), f.get("ten", "").strip(), f.get("quy_cach", "").strip(),
             f.get("dvt", "").strip(), float(f.get("gia_ban") or 0), float(f.get("gia_mua") or 0),
             f.get("ghi_chu", "").strip(), product_id),
        )
        db.commit()
        flash("Đã cập nhật sản phẩm.", "success")
        return redirect(url_for("product_list"))
    return render_template("product_form.html", product=product)


@app.route("/products/<int:product_id>/delete", methods=["POST"])
@login_required
def product_delete(product_id):
    db = get_db()
    db.execute("DELETE FROM products WHERE id=?", (product_id,))
    db.commit()
    flash("Đã xóa sản phẩm.", "success")
    return redirect(url_for("product_list"))


def extract_products_from_excel(file_stream):
    """Đọc file Excel bất kỳ, tự tìm dòng tiêu đề và nhận diện cột theo từ khóa tiếng Việt.
    Nhận diện cả trường hợp có 2 cột giá riêng (Giá bán / Giá mua) lẫn 1 cột giá chung (Đơn giá).
    Trả về list dict {ten, quy_cach, dvt, gia_ban, gia_mua} — gia_ban/gia_mua là None nếu
    cột đó không có trong file (để phân biệt với "có cột nhưng giá trị bằng 0")."""
    wb = load_workbook(file_stream, data_only=True)
    results = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row_idx = None
        col_map = {}

        # Quét tối đa 20 dòng đầu để tìm dòng tiêu đề (có ô chứa "tên" + "hàng")
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True), start=1):
            for col_idx, val in enumerate(row):
                if not val or not isinstance(val, str):
                    continue
                low = val.lower()
                if "tên" in low and ("hàng" in low or "sản phẩm" in low):
                    header_row_idx = row_idx
                    break
            if header_row_idx:
                header_values = list(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))[0]
                for col_idx, val in enumerate(header_values):
                    if not val or not isinstance(val, str):
                        continue
                    low = val.lower()
                    if "tên" in low and ("hàng" in low or "sản phẩm" in low):
                        col_map["ten"] = col_idx
                    elif "quy cách" in low:
                        col_map["quy_cach"] = col_idx
                    elif "giá bán" in low:
                        col_map["gia_ban"] = col_idx
                    elif "giá mua" in low:
                        col_map["gia_mua"] = col_idx
                    elif "đơn giá" in low or low.strip() == "giá":
                        col_map["gia_chung"] = col_idx
                    elif "đơn vị" in low or "đvt" in low:
                        col_map["dvt"] = col_idx
                break

        if not header_row_idx or "ten" not in col_map:
            continue  # sheet này không phải danh mục sản phẩm, bỏ qua

        def _to_float(raw):
            try:
                return float(raw) if raw else 0
            except (ValueError, TypeError):
                return 0

        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if col_map["ten"] >= len(row):
                continue
            ten = row[col_map["ten"]]
            if not ten or not isinstance(ten, str):
                continue
            quy_cach = row[col_map["quy_cach"]] if "quy_cach" in col_map and col_map["quy_cach"] < len(row) else ""
            dvt = row[col_map["dvt"]] if "dvt" in col_map and col_map["dvt"] < len(row) else ""

            gia_ban = _to_float(row[col_map["gia_ban"]]) if "gia_ban" in col_map and col_map["gia_ban"] < len(row) else None
            gia_mua = _to_float(row[col_map["gia_mua"]]) if "gia_mua" in col_map and col_map["gia_mua"] < len(row) else None
            gia_chung = _to_float(row[col_map["gia_chung"]]) if "gia_chung" in col_map and col_map["gia_chung"] < len(row) else None

            # Bỏ qua các dòng "rác" như nhãn chữ ký/tổng cộng vô tình rơi đúng cột tên hàng
            # (một dòng sản phẩm thật luôn có ít nhất 1 trong: quy cách/ĐVT/giá)
            if not quy_cach and not dvt and not gia_ban and not gia_mua and not gia_chung:
                continue
            if ten.strip().lower() in ("khách hàng", "người bán", "tổng cộng", "thành tiền", "ghi chú"):
                continue

            results.append({
                "ten": str(ten).strip(),
                "quy_cach": str(quy_cach).strip() if quy_cach else "",
                "dvt": str(dvt).strip() if dvt else "",
                "gia_ban": gia_ban,
                "gia_mua": gia_mua,
                "gia_chung": gia_chung,
            })

    return results


@app.route("/products/import_excel", methods=["POST"])
@login_required
def product_import_excel():
    default_nhom = request.form.get("default_nhom", "").strip()
    loai_gia = request.form.get("loai_gia", "gia_ban")  # 'gia_ban' hoặc 'gia_mua'
    try:
        markup_percent = float(request.form.get("markup_percent") or 0)
    except ValueError:
        markup_percent = 0
    file = request.files.get("excel_file")

    if not file or not file.filename:
        flash("Vui lòng chọn file Excel.", "danger")
        return redirect(url_for("product_bulk"))

    try:
        rows = extract_products_from_excel(file.stream)
    except Exception as e:
        flash(f"Không đọc được file Excel: {e}", "danger")
        return redirect(url_for("product_bulk"))

    if not rows:
        flash("Không tìm thấy cột 'Tên hàng hóa' trong file. Hãy đảm bảo file có dòng tiêu đề rõ ràng, "
              "hoặc dùng cách dán thủ công bên dưới.", "danger")
        return redirect(url_for("product_bulk"))

    db = get_db()
    created, updated, both_found = 0, 0, 0
    for r in rows:
        # Ưu tiên cột giá RIÊNG nếu file có ("Giá bán"/"Giá mua" tách biệt).
        # Chỉ dùng cột giá CHUNG ("Đơn giá") + lựa chọn loai_gia khi file không có cột riêng.
        row_gia_ban = r["gia_ban"]
        row_gia_mua = r["gia_mua"]
        if row_gia_ban is None and row_gia_mua is None and r["gia_chung"] is not None:
            if loai_gia == "gia_ban":
                row_gia_ban = r["gia_chung"]
            else:
                row_gia_mua = r["gia_chung"]

        if row_gia_ban is not None and row_gia_mua is not None:
            both_found += 1

        # Tự tính giá còn thiếu theo % lãi mong muốn (2 chiều)
        if markup_percent:
            if row_gia_ban is None and row_gia_mua is not None:
                row_gia_ban = round(row_gia_mua * (1 + markup_percent / 100))
            elif row_gia_mua is None and row_gia_ban is not None:
                row_gia_mua = round(row_gia_ban / (1 + markup_percent / 100))

        existing = db.execute(
            "SELECT * FROM products WHERE ten = ? AND IFNULL(nhom,'') = ?",
            (r["ten"], default_nhom),
        ).fetchone()
        if existing:
            # Chỉ ghi giá nếu file THỰC SỰ có giá trị cho cột đó và ô hiện tại đang trống
            # (không ghi đè giá đã nhập tay từ trước)
            if row_gia_ban is not None and not existing["gia_ban"]:
                db.execute("UPDATE products SET gia_ban = ? WHERE id = ?", (row_gia_ban, existing["id"]))
            if row_gia_mua is not None and not existing["gia_mua"]:
                db.execute("UPDATE products SET gia_mua = ? WHERE id = ?", (row_gia_mua, existing["id"]))
            if not existing["quy_cach"] and r["quy_cach"]:
                db.execute("UPDATE products SET quy_cach = ? WHERE id = ?", (r["quy_cach"], existing["id"]))
            if not existing["dvt"] and r["dvt"]:
                db.execute("UPDATE products SET dvt = ? WHERE id = ?", (r["dvt"], existing["id"]))
            updated += 1
        else:
            db.execute(
                "INSERT INTO products (nhom, ten, quy_cach, dvt, gia_ban, gia_mua) VALUES (?,?,?,?,?,?)",
                (default_nhom, r["ten"], r["quy_cach"], r["dvt"], row_gia_ban or 0, row_gia_mua or 0),
            )
            created += 1
    db.commit()

    if both_found == len(rows) and both_found > 0:
        msg = f"Đã thêm {created} sản phẩm mới, cập nhật {updated} sản phẩm đã có — đọc được cả Giá bán và Giá mua từ file."
    else:
        msg = f"Đã thêm {created} sản phẩm mới, cập nhật {updated} sản phẩm đã có."
        if markup_percent:
            msg += " Đã tự tính giá còn thiếu (nếu có) theo % lãi đã nhập."
    flash(msg, "success")
    return redirect(url_for("product_list"))


@app.route("/products/bulk", methods=["GET", "POST"])
@login_required
def product_bulk():
    if request.method == "POST":
        raw = request.form.get("bulk_text", "")
        default_nhom = request.form.get("default_nhom", "").strip()
        try:
            markup_percent = float(request.form.get("markup_percent") or 0)
        except ValueError:
            markup_percent = 0
        lines = [l for l in raw.replace("\r\n", "\n").split("\n") if l.strip()]

        db = get_db()
        created = 0
        errors = []
        for idx, line in enumerate(lines, start=1):
            cols = [c.strip() for c in line.split("\t")]
            if not cols or not cols[0]:
                errors.append(f"Dòng {idx}: thiếu tên hàng")
                continue
            try:
                ten = cols[0]
                quy_cach = cols[1] if len(cols) > 1 else ""
                dvt = cols[2] if len(cols) > 2 else ""
                gia_ban = float(re.sub(r"[^\d.\-]", "", (cols[3] if len(cols) > 3 else "0").replace(",", "")) or 0)
                gia_mua = float(re.sub(r"[^\d.\-]", "", (cols[4] if len(cols) > 4 else "0").replace(",", "")) or 0)
                ghi_chu = cols[5] if len(cols) > 5 else ""

                # Tự tính giá còn thiếu theo % lãi mong muốn (2 chiều)
                if not gia_ban and gia_mua and markup_percent:
                    gia_ban = round(gia_mua * (1 + markup_percent / 100))
                elif not gia_mua and gia_ban and markup_percent:
                    gia_mua = round(gia_ban / (1 + markup_percent / 100))

                db.execute(
                    """INSERT INTO products (nhom, ten, quy_cach, dvt, gia_ban, gia_mua, ghi_chu)
                       VALUES (?,?,?,?,?,?,?)""",
                    (default_nhom, ten, quy_cach, dvt, gia_ban, gia_mua, ghi_chu),
                )
                created += 1
            except Exception as e:
                errors.append(f"Dòng {idx}: lỗi xử lý ({e})")

        db.commit()
        if created:
            flash(f"Đã thêm {created} sản phẩm vào danh mục.", "success")
        if errors:
            flash("Một số dòng bị bỏ qua: " + " | ".join(errors[:10]) +
                  (f" (và {len(errors)-10} lỗi khác)" if len(errors) > 10 else ""), "danger")
        return redirect(url_for("product_list"))

    return render_template("product_bulk.html")


# ---------- Nhập hàng loạt đối tác (từ file Excel hoặc dán tay) ----------

def extract_partners_from_excel(file_stream):
    """Đọc file Excel bất kỳ, tự tìm cột Khách hàng/Đối tác/Nhà cung cấp + Địa chỉ/SĐT/MST.
    Xử lý luôn trường hợp tên và địa chỉ gộp chung 1 ô nhiều dòng kiểu
    'Khách hàng: Tên\\nĐịa chỉ: ...' như trong file mẫu VLXD."""
    wb = load_workbook(file_stream, data_only=True)
    results = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row_idx = None
        col_map = {}

        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True), start=1):
            for val in row:
                if not val or not isinstance(val, str):
                    continue
                low = val.lower()
                if "khách hàng" in low or "đối tác" in low or "nhà cung cấp" in low or low.strip() == "tên":
                    header_row_idx = row_idx
                    break
            if header_row_idx:
                header_values = list(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))[0]
                for col_idx, val in enumerate(header_values):
                    if not val or not isinstance(val, str):
                        continue
                    low = val.lower()
                    if "khách hàng" in low or "đối tác" in low or "nhà cung cấp" in low or low.strip() == "tên":
                        col_map["ten"] = col_idx
                    elif "địa chỉ" in low:
                        col_map["dia_chi"] = col_idx
                    elif "điện thoại" in low or "sđt" in low or "sdt" in low:
                        col_map["sdt"] = col_idx
                    elif "mã số thuế" in low or "mst" in low:
                        col_map["mst"] = col_idx
                break

        if not header_row_idx or "ten" not in col_map:
            continue

        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if col_map["ten"] >= len(row):
                continue
            raw_ten = row[col_map["ten"]]
            if not raw_ten or not isinstance(raw_ten, str):
                continue

            dia_chi = row[col_map["dia_chi"]] if "dia_chi" in col_map and col_map["dia_chi"] < len(row) else ""
            sdt = row[col_map["sdt"]] if "sdt" in col_map and col_map["sdt"] < len(row) else ""
            mst = row[col_map["mst"]] if "mst" in col_map and col_map["mst"] < len(row) else ""

            # Tách trường hợp tên gộp nhiều dòng kiểu "Khách hàng: Tên\nĐịa chỉ: ..."
            lines = [l.strip() for l in str(raw_ten).split("\n") if l.strip()]
            if not lines:
                continue
            ten_clean = re.sub(r'^(khách hàng|đối tác|nhà cung cấp)\s*:\s*', '', lines[0], flags=re.I).strip()
            for extra_line in lines[1:]:
                if extra_line.lower().startswith("địa chỉ") and not dia_chi:
                    dia_chi = extra_line

            if isinstance(dia_chi, str):
                dia_chi = re.sub(r'^địa chỉ\s*:\s*', '', dia_chi, flags=re.I).strip()
            else:
                dia_chi = ""

            if not ten_clean:
                continue
            ten_norm = ten_clean.strip().lower().rstrip(":").strip()
            skip_words = ("khách hàng", "đối tác", "nhà cung cấp", "tên", "stt",
                          "tổng cộng", "thành tiền", "ghi chú", "người bán", "người mua")
            if ten_norm in skip_words or ten_norm.startswith("địa chỉ"):
                continue

            results.append({
                "ten": ten_clean,
                "dia_chi": dia_chi,
                "sdt": str(sdt).strip() if sdt else "",
                "mst": str(mst).strip() if mst else "",
            })

    return results


@app.route("/partners/bulk", methods=["GET", "POST"])
@login_required
def partner_bulk():
    if request.method == "POST":
        raw = request.form.get("bulk_text", "")
        default_loai = request.form.get("default_loai", "buyer")
        default_nhom = request.form.get("default_nhom", "").strip()
        lines = [l for l in raw.replace("\r\n", "\n").split("\n") if l.strip()]

        db = get_db()
        created = 0
        errors = []
        for idx, line in enumerate(lines, start=1):
            cols = [c.strip() for c in line.split("\t")]
            if not cols or not cols[0]:
                errors.append(f"Dòng {idx}: thiếu tên")
                continue
            try:
                ten = cols[0]
                dia_chi = cols[1] if len(cols) > 1 else ""
                sdt = cols[2] if len(cols) > 2 else ""
                mst = cols[3] if len(cols) > 3 else ""
                ghi_chu = cols[4] if len(cols) > 4 else ""
                db.execute(
                    """INSERT INTO partners (loai, nhom, ten, dia_chi, sdt, mst, ghi_chu)
                       VALUES (?,?,?,?,?,?,?)""",
                    (default_loai, default_nhom, ten, dia_chi, sdt, mst, ghi_chu),
                )
                created += 1
            except Exception as e:
                errors.append(f"Dòng {idx}: lỗi xử lý ({e})")

        db.commit()
        if created:
            flash(f"Đã thêm {created} đối tác.", "success")
        if errors:
            flash("Một số dòng bị bỏ qua: " + " | ".join(errors[:10]), "danger")
        return redirect(url_for("partner_list"))

    return render_template("partner_bulk.html")


@app.route("/partners/import_excel", methods=["POST"])
@login_required
def partner_import_excel():
    default_nhom = request.form.get("default_nhom", "").strip()
    loai = request.form.get("loai", "buyer")
    file = request.files.get("excel_file")

    if not file or not file.filename:
        flash("Vui lòng chọn file Excel.", "danger")
        return redirect(url_for("partner_bulk"))

    try:
        rows = extract_partners_from_excel(file.stream)
    except Exception as e:
        flash(f"Không đọc được file Excel: {e}", "danger")
        return redirect(url_for("partner_bulk"))

    if not rows:
        flash("Không tìm thấy cột 'Khách hàng'/'Đối tác' trong file. Hãy dùng cách dán thủ công bên dưới.", "danger")
        return redirect(url_for("partner_bulk"))

    db = get_db()
    created, skipped = 0, 0
    for r in rows:
        existing = db.execute(
            "SELECT id FROM partners WHERE loai=? AND ten=? AND IFNULL(nhom,'')=?",
            (loai, r["ten"], default_nhom),
        ).fetchone()
        if existing:
            skipped += 1
            continue
        db.execute(
            "INSERT INTO partners (loai, nhom, ten, dia_chi, sdt, mst) VALUES (?,?,?,?,?,?)",
            (loai, default_nhom, r["ten"], r["dia_chi"], r["sdt"], r["mst"]),
        )
        created += 1
    db.commit()

    flash(f"Đã thêm {created} đối tác mới" + (f", bỏ qua {skipped} đối tác đã có sẵn." if skipped else "."), "success")
    return redirect(url_for("partner_list"))


# ---------- Tạo hóa đơn bán hàng (in được, nhiều mặt hàng) ----------

@app.route("/sales")
@login_required
def sales_list():
    nhom = request.args.get("nhom", "")
    db = get_db()
    sql = """SELECT si.*, u.full_name AS nguoi_tao_ten, u.username AS nguoi_tao_username
             FROM sales_invoices si LEFT JOIN users u ON si.nguoi_tao_id = u.id WHERE 1=1"""
    params = []
    if nhom:
        sql += " AND si.nhom = ?"
        params.append(nhom)
    sql += " ORDER BY si.ngay_lap DESC, si.id DESC"
    invoices = db.execute(sql, params).fetchall()
    nhoms = db.execute(
        "SELECT DISTINCT nhom FROM sales_invoices WHERE nhom IS NOT NULL AND nhom != '' ORDER BY nhom"
    ).fetchall()
    return render_template("sales_list.html", invoices=invoices, nhom=nhom, nhoms=nhoms)


@app.route("/sales/new", methods=["GET", "POST"])
@login_required
def sales_new():
    db = get_db()
    if request.method == "POST":
        f = request.form
        ten_hang_list = f.getlist("ten_hang[]")
        quy_cach_list = f.getlist("quy_cach[]")
        dvt_list = f.getlist("dvt[]")
        so_luong_list = f.getlist("so_luong[]")
        don_gia_list = f.getlist("don_gia[]")

        items = []
        tong_cong = 0
        stt = 0
        for i in range(len(ten_hang_list)):
            ten = ten_hang_list[i].strip()
            if not ten:
                continue
            stt += 1
            so_luong = float(so_luong_list[i] or 0)
            don_gia = float(don_gia_list[i] or 0)
            thanh_tien = so_luong * don_gia
            tong_cong += thanh_tien
            items.append({
                "stt": stt,
                "ten_hang": ten,
                "quy_cach": quy_cach_list[i].strip(),
                "dvt": dvt_list[i].strip(),
                "so_luong": so_luong,
                "don_gia": don_gia,
                "thanh_tien": thanh_tien,
            })

        partners_json = json.dumps([dict(p) for p in db.execute("SELECT * FROM partners ORDER BY nhom, loai, ten").fetchall()])
        products_json = json.dumps([dict(p) for p in db.execute("SELECT * FROM products ORDER BY nhom, ten").fetchall()])

        if not items:
            flash("Cần ít nhất 1 mặt hàng.", "danger")
            return render_template(
                "sales_form.html", invoice=None, items=[],
                today=datetime.now().strftime("%Y-%m-%d"), partners_json=partners_json,
                products_json=products_json,
            )

        nhom = f.get("nhom", "").strip()
        seller_name = f.get("seller_name", "").strip()
        seller_slogan = f.get("seller_slogan", "").strip()
        buyer_name = f.get("buyer_name", "").strip()
        buyer_address = f.get("buyer_address", "").strip()
        buyer_phone = f.get("buyer_phone", "").strip()
        dia_diem = f.get("dia_diem", "").strip()
        loai_chung_tu = f.get("loai_chung_tu", "hoa_don")
        try:
            thanh_toan = float(f.get("thanh_toan") or 0)
        except ValueError:
            thanh_toan = 0

        cur = db.execute(
            """INSERT INTO sales_invoices
               (nhom, so_hd, ngay_lap, seller_name, seller_slogan,
                buyer_name, buyer_address, buyer_phone, dia_diem, tong_cong,
                loai_chung_tu, thanh_toan, nguoi_tao_id)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                nhom, f.get("so_hd", "").strip(), f.get("ngay_lap"),
                seller_name, seller_slogan, buyer_name,
                buyer_address, buyer_phone, dia_diem, tong_cong,
                loai_chung_tu, thanh_toan, current_user.id,
            ),
        )
        invoice_id = cur.lastrowid
        for it in items:
            db.execute(
                """INSERT INTO sales_invoice_items
                   (invoice_id, stt, ten_hang, quy_cach, dvt, so_luong, don_gia, thanh_tien)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (invoice_id, it["stt"], it["ten_hang"], it["quy_cach"], it["dvt"],
                 it["so_luong"], it["don_gia"], it["thanh_tien"]),
            )

        # Lưu đối tác mới vào danh mục nếu được tick và chưa tồn tại
        if f.get("save_seller") and seller_name:
            exists = db.execute(
                "SELECT id FROM partners WHERE loai='seller' AND ten=? AND IFNULL(nhom,'')=?",
                (seller_name, nhom),
            ).fetchone()
            if not exists:
                db.execute(
                    "INSERT INTO partners (loai, nhom, ten) VALUES ('seller',?,?)",
                    (nhom, seller_name),
                )
        if f.get("save_buyer") and buyer_name:
            exists = db.execute(
                "SELECT id FROM partners WHERE loai='buyer' AND ten=? AND IFNULL(nhom,'')=?",
                (buyer_name, nhom),
            ).fetchone()
            if not exists:
                db.execute(
                    "INSERT INTO partners (loai, nhom, ten, dia_chi, sdt) VALUES ('buyer',?,?,?,?)",
                    (nhom, buyer_name, buyer_address, buyer_phone),
                )
        db.commit()

        flash("Đã tạo hóa đơn bán lẻ.", "success")
        return redirect(url_for("sales_view", invoice_id=invoice_id))

    partners_json = json.dumps([dict(p) for p in db.execute("SELECT * FROM partners ORDER BY nhom, loai, ten").fetchall()])
    products_json = json.dumps([dict(p) for p in db.execute("SELECT * FROM products ORDER BY nhom, ten").fetchall()])
    return render_template(
        "sales_form.html", invoice=None, items=[],
        today=datetime.now().strftime("%Y-%m-%d"), partners_json=partners_json,
        products_json=products_json,
    )


@app.route("/sales/<int:invoice_id>")
@login_required
def sales_view(invoice_id):
    db = get_db()
    invoice = db.execute("SELECT * FROM sales_invoices WHERE id=?", (invoice_id,)).fetchone()
    if not invoice:
        flash("Không tìm thấy hóa đơn.", "danger")
        return redirect(url_for("sales_list"))
    items = db.execute(
        "SELECT * FROM sales_invoice_items WHERE invoice_id=? ORDER BY stt", (invoice_id,)
    ).fetchall()
    return render_template(
        "sales_view.html", invoice=invoice, items=items,
        so_tien_chu=so_thanh_chu(invoice["tong_cong"]),
    )


@app.route("/sales/<int:invoice_id>/delete", methods=["POST"])
@login_required
def sales_delete(invoice_id):
    db = get_db()
    db.execute("DELETE FROM sales_invoice_items WHERE invoice_id=?", (invoice_id,))
    db.execute("DELETE FROM sales_invoices WHERE id=?", (invoice_id,))
    db.commit()
    flash("Đã xóa hóa đơn.", "success")
    return redirect(url_for("sales_list"))


def build_sales_invoice_pdf(invoice, items, page_size="a4"):
    is_a5 = page_size.lower() == "a5"
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A5 if is_a5 else A4,
        topMargin=(10 if is_a5 else 15) * mm, bottomMargin=(10 if is_a5 else 15) * mm,
        leftMargin=(10 if is_a5 else 15) * mm, rightMargin=(10 if is_a5 else 15) * mm,
    )

    fs = 0.8 if is_a5 else 1.0  # hệ số thu nhỏ cỡ chữ cho khổ A5

    style_normal = ParagraphStyle("normal", fontName="VNSans", fontSize=10 * fs, leading=13 * fs)
    style_bold = ParagraphStyle("bold", fontName="VNSans-Bold", fontSize=12 * fs, leading=15 * fs)
    style_slogan = ParagraphStyle("slogan", fontName="VNSans", fontSize=9 * fs, leading=12 * fs)
    style_title = ParagraphStyle(
        "title", fontName="VNSans-Bold", fontSize=16 * fs, leading=20 * fs, alignment=TA_CENTER,
        spaceAfter=4,
    )
    style_quochieu = ParagraphStyle(
        "quochieu", fontName="VNSans-Bold", fontSize=11 * fs, leading=14 * fs, alignment=TA_CENTER,
    )
    style_center = ParagraphStyle("center", fontName="VNSans", fontSize=10 * fs, alignment=TA_CENTER)
    style_right = ParagraphStyle("right", fontName="VNSans", fontSize=10 * fs, alignment=TA_RIGHT)
    style_cell = ParagraphStyle("cell", fontName="VNSans", fontSize=(8 if is_a5 else 9), leading=(10 if is_a5 else 11))
    style_cell_bold = ParagraphStyle("cell_bold", fontName="VNSans-Bold", fontSize=(8 if is_a5 else 9), leading=(10 if is_a5 else 11))
    style_cell_center = ParagraphStyle("cell_center", parent=style_cell, alignment=TA_CENTER)
    style_cell_right = ParagraphStyle("cell_right", parent=style_cell, alignment=TA_RIGHT)
    style_cell_bold_center = ParagraphStyle("cell_bold_center", parent=style_cell_bold, alignment=TA_CENTER)
    style_cell_bold_right = ParagraphStyle("cell_bold_right", parent=style_cell_bold, alignment=TA_RIGHT)
    style_red = ParagraphStyle("red", fontName="VNSans-Bold", fontSize=10 * fs, textColor=colors.HexColor("#C00000"))

    story = []

    left_w = 75 * mm if is_a5 else 85 * mm
    right_w = 55 * mm if is_a5 else 85 * mm

    left_cell = [Paragraph(f"CƠ SỞ {invoice['seller_name'] or ''}".upper(), style_bold)]
    if invoice["seller_slogan"]:
        left_cell.append(Paragraph(invoice["seller_slogan"], style_slogan))
    right_cell = [
        Paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", style_quochieu),
        Paragraph("<u>Độc lập – Tự do – Hạnh Phúc</u>", style_quochieu),
    ]
    header_table = Table([[left_cell, right_cell]], colWidths=[left_w, right_w])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 10 if is_a5 else 14))

    is_phieu = invoice["loai_chung_tu"] == "phieu_giao"

    story.append(Paragraph("PHIẾU GIAO HÀNG" if is_phieu else "HÓA ĐƠN BÁN LẺ", style_title))
    if is_phieu:
        ngay_str_top = invoice["ngay_lap"]
        try:
            d0 = datetime.strptime(ngay_str_top, "%Y-%m-%d")
            ngay_phrase_top = f"ngày {d0.day} tháng {d0.month:02d} năm {d0.year}"
        except Exception:
            ngay_phrase_top = ngay_str_top
        dia_diem_top = invoice["dia_diem"] or ""
        dong_ngay_top = f"{dia_diem_top}, {ngay_phrase_top}" if dia_diem_top else ngay_phrase_top.capitalize()
        story.append(Paragraph(f"<i>{dong_ngay_top}</i>", style_center))
    elif invoice["so_hd"]:
        story.append(Paragraph(f"Số: {invoice['so_hd']}", style_center))
    story.append(Spacer(1, 8 if is_a5 else 10))

    story.append(Paragraph(f"- Họ và tên người nhận hàng: {invoice['buyer_name']}", style_normal))
    story.append(Paragraph(f"- Địa chỉ: {invoice['buyer_address'] or ''}", style_normal))
    story.append(Paragraph(f"- Số điện thoại: {invoice['buyer_phone'] or ''}", style_normal))
    if is_phieu:
        story.append(Paragraph(
            "<i>Tôi/chúng tôi tiến hành bàn giao cho Ông/Bà hàng hóa theo bảng kê dưới đây:</i>",
            style_normal,
        ))
    story.append(Spacer(1, 4))

    dvt_para = Paragraph("<i>ĐVT: Đồng</i>", style_right)
    story.append(dvt_para)
    story.append(Spacer(1, 2))

    header = ["STT", "Tên hàng", "ĐVT", "Đơn giá", "Số lượng", "Thành tiền"]
    header_styles = [style_cell_bold_center, style_cell_bold, style_cell_bold_center,
                      style_cell_bold_right, style_cell_bold_right, style_cell_bold_right]
    table_data = [[Paragraph(h, st) for h, st in zip(header, header_styles)]]

    tong_so_luong = 0
    for it in items:
        tong_so_luong += it["so_luong"]
        sl_str = f"{it['so_luong']:,.0f}" if it["so_luong"] % 1 == 0 else f"{it['so_luong']:,.2f}"
        table_data.append([
            Paragraph(str(it["stt"]), style_cell_center),
            Paragraph(it["ten_hang"], style_cell),
            Paragraph(it["dvt"] or "", style_cell_center),
            Paragraph(f"{it['don_gia']:,.0f}", style_cell_right),
            Paragraph(sl_str, style_cell_right),
            Paragraph(f"{it['thanh_tien']:,.0f}", style_cell_right),
        ])

    tong_sl_str = f"{tong_so_luong:,.0f}" if tong_so_luong % 1 == 0 else f"{tong_so_luong:,.2f}"
    table_data.append([
        Paragraph("", style_cell), Paragraph("Tổng cộng", style_cell_bold), Paragraph("", style_cell),
        Paragraph("", style_cell),
        Paragraph(tong_sl_str, style_cell_bold_right),
        Paragraph(f"{invoice['tong_cong']:,.0f}", style_cell_bold_right),
    ])

    if is_a5:
        col_widths = [13 * mm, 35 * mm, 13 * mm, 21 * mm, 15 * mm, 25 * mm]
    else:
        col_widths = [12 * mm, 50 * mm, 18 * mm, 30 * mm, 25 * mm, 35 * mm]
    tbl = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#333333")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E5266")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 0), (2, -1), "CENTER"),
        ("ALIGN", (3, 0), (5, -1), "RIGHT"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#D6E4F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 3 if is_a5 else 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3 if is_a5 else 4),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 8 if is_a5 else 10))

    so_chu = so_thanh_chu(invoice["tong_cong"])
    footer_block = []

    if is_phieu:
        thanh_toan = float(invoice["thanh_toan"] or 0)
        con_lai = invoice["tong_cong"] - thanh_toan
        footer_block.append(Paragraph(f"- Tổng số tiền: <b>{invoice['tong_cong']:,.0f} đồng</b>", style_normal))
        footer_block.append(Paragraph(f"- Thanh toán: <b>{thanh_toan:,.0f} đồng</b>", style_normal))
        footer_block.append(Paragraph(f"- Còn lại: <b>{con_lai:,.0f} đồng</b>", style_red))
    else:
        footer_block.append(Paragraph(
            f"Thành tiền: {invoice['tong_cong']:,.0f} đồng (Bằng chữ: <i>{so_chu.rstrip('.')}</i>).",
            style_red,
        ))
    footer_block.append(Spacer(1, 6 if is_a5 else 8))

    ngay_str = invoice["ngay_lap"]
    try:
        d = datetime.strptime(ngay_str, "%Y-%m-%d")
        ngay_phrase = f"ngày {d.day} tháng {d.month} năm {d.year}"
    except Exception:
        ngay_phrase = ngay_str
    dia_diem = invoice["dia_diem"] or ""
    dòng_ngay = f"{dia_diem}, {ngay_phrase}" if dia_diem else ngay_phrase.capitalize()
    if not is_phieu:
        footer_block.append(Paragraph(dòng_ngay, style_right))
    footer_block.append(Spacer(1, 16 if is_a5 else 20))

    if is_phieu:
        sig_table = Table(
            [[Paragraph("BÊN GIAO", style_center), Paragraph("KHÁCH HÀNG", style_center)]],
            colWidths=[left_w, right_w],
        )
        footer_block.append(sig_table)
        footer_block.append(Spacer(1, 30 if is_a5 else 40))
        name_table = Table(
            [[Paragraph(invoice["seller_name"] or "", style_bold), Paragraph("", style_center)]],
            colWidths=[left_w, right_w],
        )
        name_table.setStyle(TableStyle([("ALIGN", (0, 0), (0, 0), "CENTER")]))
        footer_block.append(name_table)
    else:
        sig_table = Table(
            [[Paragraph("Khách hàng", style_center), Paragraph("Người bán hàng", style_center)]],
            colWidths=[left_w, right_w],
        )
        footer_block.append(sig_table)

    story.append(KeepTogether(footer_block))

    doc.build(story)
    buf.seek(0)
    return buf


# ---------- Word export (.docx) khổ A4/A5 ----------

def _set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_run(paragraph, text, bold=False, size=11, color=None, italic=False, underline=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def build_sales_invoice_docx(invoice, items, page_size="a4"):
    is_a5 = page_size.lower() == "a5"
    docx_doc = DocxDocument()

    section = docx_doc.sections[0]
    if is_a5:
        section.page_width = Mm(148)
        section.page_height = Mm(210)
        margin = Mm(12)
    else:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        margin = Mm(18)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin

    fs = 9 if is_a5 else 11
    fs_title = 15 if is_a5 else 18
    fs_cell = 8 if is_a5 else 10

    # Header: bảng 1 dòng 2 cột — trái tên cơ sở, phải quốc hiệu
    header_table = docx_doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_w = Mm(148 - 24) if is_a5 else Mm(210 - 36)
    left_w_docx = Mm((148 - 24) * 0.55) if is_a5 else Mm((210 - 36) * 0.55)
    right_w_docx = Mm((148 - 24) * 0.45) if is_a5 else Mm((210 - 36) * 0.45)
    header_table.columns[0].width = left_w_docx
    header_table.columns[1].width = right_w_docx
    left_cell, right_cell = header_table.rows[0].cells
    left_cell.width = left_w_docx
    right_cell.width = right_w_docx

    p1 = left_cell.paragraphs[0]
    _set_run(p1, f"CƠ SỞ {(invoice['seller_name'] or '').upper()}", bold=True, size=fs + 1)
    if invoice["seller_slogan"]:
        p1b = left_cell.add_paragraph()
        _set_run(p1b, invoice["seller_slogan"], size=fs - 1)

    p2 = right_cell.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p2, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, size=fs)
    p3 = right_cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p3, "Độc lập – Tự do – Hạnh Phúc", bold=True, size=fs, underline=True)

    docx_doc.add_paragraph()

    is_phieu = invoice["loai_chung_tu"] == "phieu_giao"

    p_title = docx_doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p_title, "PHIẾU GIAO HÀNG" if is_phieu else "HÓA ĐƠN BÁN LẺ", bold=True, size=fs_title)

    ngay_str_top = invoice["ngay_lap"]
    try:
        d0 = datetime.strptime(ngay_str_top, "%Y-%m-%d")
        ngay_phrase_top = f"ngày {d0.day} tháng {d0.month:02d} năm {d0.year}"
    except Exception:
        ngay_phrase_top = ngay_str_top
    dia_diem_top = invoice["dia_diem"] or ""
    dong_ngay_top = f"{dia_diem_top}, {ngay_phrase_top}" if dia_diem_top else ngay_phrase_top.capitalize()

    if is_phieu:
        p_so = docx_doc.add_paragraph()
        p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p_so, dong_ngay_top, italic=True, size=fs)
    elif invoice["so_hd"]:
        p_so = docx_doc.add_paragraph()
        p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p_so, f"Số: {invoice['so_hd']}", size=fs)

    docx_doc.add_paragraph()

    for line in (
        f"- Họ và tên người nhận hàng: {invoice['buyer_name']}",
        f"- Địa chỉ: {invoice['buyer_address'] or ''}",
        f"- Số điện thoại: {invoice['buyer_phone'] or ''}",
    ):
        pb = docx_doc.add_paragraph()
        _set_run(pb, line, size=fs)

    if is_phieu:
        p_transition = docx_doc.add_paragraph()
        _set_run(p_transition, "Tôi/chúng tôi tiến hành bàn giao cho Ông/Bà hàng hóa theo bảng kê dưới đây:",
                  italic=True, size=fs)

    p_dvt = docx_doc.add_paragraph()
    p_dvt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run(p_dvt, "ĐVT: Đồng", italic=True, size=fs - 1)

    headers = ["STT", "Tên hàng", "ĐVT", "Đơn giá", "Số lượng", "Thành tiền"]
    table = docx_doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Cho dòng tiêu đề tự lặp lại ở đầu mỗi trang khi bảng dài nhiều trang
    header_row = table.rows[0]
    trPr = header_row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(cell.paragraphs[0], h, bold=True, size=fs_cell, color="FFFFFF")
        _set_cell_shading(cell, "2E5266")

    tong_so_luong = 0
    for it in items:
        tong_so_luong += it["so_luong"]
        sl_str = f"{it['so_luong']:,.0f}" if it["so_luong"] % 1 == 0 else f"{it['so_luong']:,.2f}"
        row = table.add_row().cells
        values = [str(it["stt"]), it["ten_hang"], it["dvt"] or "",
                  f"{it['don_gia']:,.0f}", sl_str, f"{it['thanh_tien']:,.0f}"]
        aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
        for i, val in enumerate(values):
            row[i].paragraphs[0].alignment = aligns[i]
            _set_run(row[i].paragraphs[0], val, size=fs_cell)

    tong_sl_str = f"{tong_so_luong:,.0f}" if tong_so_luong % 1 == 0 else f"{tong_so_luong:,.2f}"
    total_row = table.add_row().cells
    total_values = ["", "Tổng cộng", "", "", tong_sl_str, f"{invoice['tong_cong']:,.0f}"]
    total_aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                     WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
    for i, val in enumerate(total_values):
        total_row[i].paragraphs[0].alignment = total_aligns[i]
        _set_run(total_row[i].paragraphs[0], val, bold=True, size=fs_cell)
        _set_cell_shading(total_row[i], "D6E4F0")

    docx_doc.add_paragraph()
    so_chu = so_thanh_chu(invoice["tong_cong"])

    if is_phieu:
        thanh_toan = float(invoice["thanh_toan"] or 0)
        con_lai = invoice["tong_cong"] - thanh_toan
        p_tong = docx_doc.add_paragraph()
        p_tong.paragraph_format.keep_with_next = True
        _set_run(p_tong, f"- Tổng số tiền: {invoice['tong_cong']:,.0f} đồng", bold=True, size=fs)
        p_ttoan = docx_doc.add_paragraph()
        p_ttoan.paragraph_format.keep_with_next = True
        _set_run(p_ttoan, f"- Thanh toán: {thanh_toan:,.0f} đồng", bold=True, size=fs)
        p_conlai = docx_doc.add_paragraph()
        p_conlai.paragraph_format.keep_with_next = True
        _set_run(p_conlai, f"- Còn lại: {con_lai:,.0f} đồng", bold=True, size=fs, color="C00000")
    else:
        p_total = docx_doc.add_paragraph()
        p_total.paragraph_format.keep_with_next = True
        _set_run(
            p_total,
            f"Thành tiền: {invoice['tong_cong']:,.0f} đồng (Bằng chữ: {so_chu.rstrip('.')}).",
            bold=True, size=fs, color="C00000",
        )

    ngay_str = invoice["ngay_lap"]
    try:
        d = datetime.strptime(ngay_str, "%Y-%m-%d")
        ngay_phrase = f"ngày {d.day} tháng {d.month} năm {d.year}"
    except Exception:
        ngay_phrase = ngay_str
    dia_diem = invoice["dia_diem"] or ""
    dòng_ngay = f"{dia_diem}, {ngay_phrase}" if dia_diem else ngay_phrase.capitalize()
    if not is_phieu:
        p_ngay = docx_doc.add_paragraph()
        p_ngay.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ngay.paragraph_format.keep_with_next = True
        _set_run(p_ngay, dòng_ngay, size=fs)

    p_spacer1 = docx_doc.add_paragraph()
    p_spacer1.paragraph_format.keep_with_next = True
    p_spacer2 = docx_doc.add_paragraph()
    p_spacer2.paragraph_format.keep_with_next = True

    sig_table = docx_doc.add_table(rows=1, cols=2)
    sig_table.autofit = False
    sig_table.columns[0].width = left_w_docx
    sig_table.columns[1].width = right_w_docx
    left_sig, right_sig = sig_table.rows[0].cells
    left_sig.width = left_w_docx
    right_sig.width = right_w_docx
    left_sig.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_sig.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if is_phieu:
        _set_run(left_sig.paragraphs[0], "BÊN GIAO", bold=True, size=fs)
        _set_run(right_sig.paragraphs[0], "KHÁCH HÀNG", bold=True, size=fs)
        for _ in range(3):
            docx_doc.add_paragraph()
        p_name = docx_doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_run(p_name, invoice["seller_name"] or "", bold=True, size=fs)
    else:
        _set_run(left_sig.paragraphs[0], "Khách hàng", bold=True, size=fs)
        _set_run(right_sig.paragraphs[0], "Người bán hàng", bold=True, size=fs)

    buf = BytesIO()
    docx_doc.save(buf)
    buf.seek(0)
    return buf


@app.route("/sales/<int:invoice_id>/pdf")
@login_required
def sales_pdf(invoice_id):
    size = request.args.get("size", "a4")
    db = get_db()
    invoice = db.execute("SELECT * FROM sales_invoices WHERE id=?", (invoice_id,)).fetchone()
    if not invoice:
        flash("Không tìm thấy hóa đơn.", "danger")
        return redirect(url_for("sales_list"))
    items = db.execute(
        "SELECT * FROM sales_invoice_items WHERE invoice_id=? ORDER BY stt", (invoice_id,)
    ).fetchall()
    buf = build_sales_invoice_pdf(invoice, items, page_size=size)
    filename = f"HoaDon_{invoice['so_hd'] or invoice_id}_{size.upper()}.pdf"
    return send_file(buf, as_attachment=True, download_name=filename, mimetype="application/pdf")


@app.route("/sales/<int:invoice_id>/docx")
@login_required
def sales_docx(invoice_id):
    size = request.args.get("size", "a4")
    db = get_db()
    invoice = db.execute("SELECT * FROM sales_invoices WHERE id=?", (invoice_id,)).fetchone()
    if not invoice:
        flash("Không tìm thấy hóa đơn.", "danger")
        return redirect(url_for("sales_list"))
    items = db.execute(
        "SELECT * FROM sales_invoice_items WHERE invoice_id=? ORDER BY stt", (invoice_id,)
    ).fetchall()
    buf = build_sales_invoice_docx(invoice, items, page_size=size)
    filename = f"HoaDon_{invoice['so_hd'] or invoice_id}_{size.upper()}.docx"
    return send_file(
        buf, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
else:
    # Ensure DB exists when run under gunicorn (Railway)
    init_db()
